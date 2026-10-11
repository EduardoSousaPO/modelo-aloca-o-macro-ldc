"""F-031 — testes da montagem do pacote de entrega ao Comitê.

**RF-043 / CA-113, CA-114 · RF-044 / CA-115, CA-116 · RF-045 / CA-117 · CA-121 (idempotência)
· CA-090 como gate.**

| Bloco | O que trava | CA |
|---|---|---|
| 1 | **lista de inclusão × exclusão**: o manifesto vence a lista,
      `publicavel: false` ⇒ `BLOQUEADO` | CA-113, D-af |
| 2 | **CSV pt-BR**: `;`, vírgula decimal, UTF-8 **com BOM**, `LF` — e a volta lê o mesmo | CA-113 |
| 3 | **zip determinístico**: timestamps fixos, ordem fixa,
      duas montagens ⇒ mesmo `sha256` | CA-121 |
| 4 | **smoke** detecta artefato proibido injetado, hash divergente e arquivo ausente | CA-116 |
| 5 | **gate lido do disco**: o selo nasce do `ADR-001` em `tmp_path`,
      nunca de texto digitado | CA-090, CA-115 |
| 6 | os textos lidos do disco: as **10** limitações, as **12** decisões,
      as substituições | CA-117 |
| 7 | **acervo real** `data/canonical/`
      (padrão `Acervo`; desligado por `LDC_IGNORAR_DATA=1`) | todos |

Os blocos 1-6 rodam sobre **fixtures sintéticas** montadas em `tmp_path` e nomeadas como tal
(Anti-SPEC §10.1.2) — valores redondos de propósito (`+5%`, `-2%`), para que confundir sintético
com acervo seja impossível na leitura. **Nenhum teste escreve em `data/canonical/` do repositório.**
"""

from __future__ import annotations

import json
import os
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd
import pytest

from apps.analysis.entrega import montar_pacote as P
from apps.pipeline.common.errors import SerieInconsistente
from apps.pipeline.common.io_layers import caminho_canonical

COMO_GERAR = "rode `uv run python -m apps.analysis.entrega.montar_pacote --smoke`"
VAR_IGNORAR_DATA = "LDC_IGNORAR_DATA"

SHA_SINTETICO = "f" * 64
"""Hash obviamente falso: nenhum artefato real tem `sha256` de um só dígito repetido."""


# ═════════════════════════════════════════════════════════════════════════════
# Fixtura sintética: um `MANIFEST.json` mínimo
# ═════════════════════════════════════════════════════════════════════════════


def manifesto_sintetico(*, publicavel_de: str | None = None, linhas: int = 83) -> dict[str, Any]:
    """Um manifesto de mentirinha com os campos que `montar_pacote` consulta.

    `publicavel_de` marca **esse** artefato como `publicavel: false` — é como o bloco 1
    injeta a situação que o manifesto real tem hoje com `internacional_pesos.parquet`.
    """
    artefatos = [
        {
            "arquivo": nome,
            "bytes": 1,
            "linhas": linhas,
            "colunas": 3,
            "sha256": SHA_SINTETICO,
            "determinista": True,
            "publicavel": nome != publicavel_de,
            "coluna_de_data": "quarter",
            "date_min": "2006Q1",
            "date_max": "2026Q3",
        }
        for nome in P._arquivos_esperados_no_manifesto()
    ]
    if publicavel_de is not None:
        for entrada in artefatos:
            if entrada["arquivo"] == publicavel_de:
                entrada["motivo_nao_publicavel"] = "fixture sintetica: proibido de proposito"
    return {
        "manifesto_versao": 1,
        "feature": "F-029",
        "agregado": {"sha256": SHA_SINTETICO, "n_artefatos": len(artefatos)},
        "artefatos": artefatos,
        "data_do_dado": {
            "primeiro_trimestre_macro": "2006Q1",
            "ultimo_trimestre_macro": "2026Q3",
        },
        "convencao": {"agregado": "fixture sintetica"},
    }


@pytest.fixture
def quadro_sintetico() -> pd.DataFrame:
    """3 linhas com números redondos — `+5%`, `-2%`, `0%` — e uma coluna de texto."""
    return pd.DataFrame(
        {
            "quarter": ["2006Q1", "2006Q2", "2006Q3"],
            "asset_id": ["FAKE_A", "FAKE_B", "FAKE_C"],
            "ret_tri": [0.05, -0.02, 0.0],
        }
    )


# ═════════════════════════════════════════════════════════════════════════════
# 1 — Lista de inclusão × exclusão: o manifesto vence a lista
# ═════════════════════════════════════════════════════════════════════════════


def test_a_lista_de_inclusao_tem_as_10_tabelas_e_nao_tem_internacional_pesos() -> None:
    """`internacional_pesos` é o artefato que D-af proíbe — não pode estar na lista."""
    assert len(P.TABELAS_DO_PACOTE) == 10
    assert "internacional_pesos" not in P.TABELAS_DO_PACOTE
    assert len(set(P.TABELAS_DO_PACOTE)) == len(P.TABELAS_DO_PACOTE)


def test_a_lista_de_inclusao_nao_contem_nada_de_ARTEFATOS_NAO_PUBLICAVEIS() -> None:
    """A trava estática: nenhum nome da constante de `run.py` aparece na lista de inclusão."""
    from apps.pipeline.run import ARTEFATOS_NAO_PUBLICAVEIS

    assert set(P._arquivos_esperados_no_manifesto()).isdisjoint(ARTEFATOS_NAO_PUBLICAVEIS)


def test_tudo_que_nao_esta_na_lista_fica_fora() -> None:
    """Artefato real do canônico que o pacote **não** leva — a exclusão por omissão."""
    esperados = set(P._arquivos_esperados_no_manifesto())
    for fora in (
        "raw_series.parquet",
        "alinhado_trimestral.parquet",
        "series_catalog.parquet",
        "macro_quarterly_provenance.parquet",
        "regime_neighbors.parquet",
        "ldc.duckdb",
        "internacional_pesos.parquet",
    ):
        assert fora not in esperados


def test_nao_publicaveis_sai_do_manifesto_em_disco_nao_da_constante() -> None:
    """`nao_publicaveis` lê o campo do manifesto — a constante é só complemento do motivo."""
    manifesto = manifesto_sintetico(publicavel_de="macro_quarterly.parquet")
    proibidos = P.nao_publicaveis(manifesto)
    assert proibidos == {"macro_quarterly.parquet": "fixture sintetica: proibido de proposito"}


def test_validar_lista_de_inclusao_bloqueia_artefato_marcado_nao_publicavel() -> None:
    """CA-116: um nome da lista marcado `publicavel: false` ⇒ `BLOQUEADO`, não aviso."""
    manifesto = manifesto_sintetico(publicavel_de="regimes_v1.parquet")
    with pytest.raises(SerieInconsistente, match="ARTEFATO NÃO PUBLICÁVEL"):
        P.validar_lista_de_inclusao(manifesto)


def test_validar_lista_de_inclusao_bloqueia_artefato_ausente_do_manifesto() -> None:
    """Artefato da lista que o manifesto não conhece ⇒ `BLOQUEADO — FONTE INDISPONÍVEL`."""
    manifesto = manifesto_sintetico()
    manifesto["artefatos"] = [
        e for e in manifesto["artefatos"] if e["arquivo"] != "macro_quarterly.parquet"
    ]
    with pytest.raises(SerieInconsistente, match="FONTE INDISPONÍVEL"):
        P.validar_lista_de_inclusao(manifesto)


def test_validar_lista_de_inclusao_passa_no_manifesto_completo() -> None:
    P.validar_lista_de_inclusao(manifesto_sintetico())


def test_ler_manifesto_bloqueia_quando_o_arquivo_nao_existe(tmp_path: Path) -> None:
    (tmp_path / "data" / "canonical").mkdir(parents=True)
    with pytest.raises(SerieInconsistente, match="FONTE INDISPONÍVEL"):
        P.ler_manifesto(raiz=tmp_path)


# ═════════════════════════════════════════════════════════════════════════════
# 2 — CSV pt-BR
# ═════════════════════════════════════════════════════════════════════════════


def test_csv_ptbr_usa_ponto_e_virgula_virgula_decimal_e_bom(
    tmp_path: Path, quadro_sintetico: pd.DataFrame
) -> None:
    """CA-113: o CSV abre no Excel pt-BR sem assistente — BOM, `;` e vírgula decimal."""
    alvo = P.gravar_csv_ptbr(quadro_sintetico, tmp_path / "fake.csv")
    bruto = alvo.read_bytes()
    assert bruto.startswith(b"\xef\xbb\xbf"), "falta o BOM de UTF-8"
    texto = bruto.decode("utf-8-sig")
    assert texto.splitlines()[0] == "quarter;asset_id;ret_tri"
    assert "0,05" in texto
    assert "-0,02" in texto
    assert "0.05" not in texto


def test_csv_ptbr_usa_lf_e_nao_crlf(tmp_path: Path, quadro_sintetico: pd.DataFrame) -> None:
    """CRLF mudaria o `sha256` sem mudar o conteúdo — é a armadilha D-v, no Windows."""
    alvo = P.gravar_csv_ptbr(quadro_sintetico, tmp_path / "fake.csv")
    assert b"\r\n" not in alvo.read_bytes()


def test_csv_ptbr_ida_e_volta_preserva_valores(
    tmp_path: Path, quadro_sintetico: pd.DataFrame
) -> None:
    """A volta é a que o smoke usa para **abrir** cada CSV: tem de reconstruir o quadro."""
    alvo = P.gravar_csv_ptbr(quadro_sintetico, tmp_path / "fake.csv")
    pd.testing.assert_frame_equal(P.ler_csv_ptbr(alvo), quadro_sintetico)


def test_csv_ptbr_e_deterministico(tmp_path: Path, quadro_sintetico: pd.DataFrame) -> None:
    primeiro = P.gravar_csv_ptbr(quadro_sintetico, tmp_path / "a.csv").read_bytes()
    segundo = P.gravar_csv_ptbr(quadro_sintetico, tmp_path / "b.csv").read_bytes()
    assert primeiro == segundo


def test_a_convencao_de_csv_de_f024_e_outra_e_esta_declarada(tmp_path: Path) -> None:
    """As tabelas de 1 página vêm com `,` e sem BOM — a diferença é declarada, não escondida."""
    assert P.CSV_DE_F024_SEPARADOR != P.CSV_SEPARADOR
    assert P.CSV_DE_F024_ENCODING == "utf-8"
    _escrever_adr(tmp_path, "Proposta")
    leia_me = " ".join(
        str(valor)
        for linha in P._linhas_do_leia_me(
            gate=P.ler_gate(raiz=tmp_path), manifesto=manifesto_sintetico()
        )
        for valor in linha
    )
    assert "FORMATO DOS CSV EM dados/" in leia_me
    assert "FORMATO DOS CSV EM tabelas/" in leia_me
    assert "convencao DIFERENTE" in leia_me


# ═════════════════════════════════════════════════════════════════════════════
# 3 — Zip determinístico
# ═════════════════════════════════════════════════════════════════════════════


def _arvore_sintetica(raiz: Path) -> Path:
    pacote = raiz / "pacote"
    (pacote / "dados").mkdir(parents=True)
    (pacote / "b.txt").write_text("b", encoding="utf-8")
    (pacote / "a.txt").write_text("a", encoding="utf-8")
    (pacote / "dados" / "c.txt").write_text("c", encoding="utf-8")
    return pacote


def test_zip_tem_timestamp_fixo_em_toda_entrada(tmp_path: Path) -> None:
    """CA-121: o relógio não entra no `.zip` — senão o hash do pacote muda a cada montagem."""
    pacote = _arvore_sintetica(tmp_path)
    P.zipar_deterministico(pacote, tmp_path / "p.zip")
    with zipfile.ZipFile(tmp_path / "p.zip") as arquivo:
        assert {info.date_time for info in arquivo.infolist()} == {P.DATA_FIXA_NO_ZIP}


def test_zip_tem_ordem_fixa_alfabetica_do_caminho_posix(tmp_path: Path) -> None:
    pacote = _arvore_sintetica(tmp_path)
    P.zipar_deterministico(pacote, tmp_path / "p.zip")
    with zipfile.ZipFile(tmp_path / "p.zip") as arquivo:
        assert [info.filename for info in arquivo.infolist()] == [
            "pacote/a.txt",
            "pacote/b.txt",
            "pacote/dados/c.txt",
        ]


def test_zip_duas_geracoes_dao_o_mesmo_sha256(tmp_path: Path) -> None:
    pacote = _arvore_sintetica(tmp_path)
    primeiro = P.zipar_deterministico(pacote, tmp_path / "p1.zip")
    segundo = P.zipar_deterministico(pacote, tmp_path / "p2.zip")
    assert primeiro == segundo
    assert primeiro == P.sha256_do_arquivo(tmp_path / "p2.zip")


def test_rezipar_fixa_o_timestamp_e_o_dcterms_modified_do_xlsx(tmp_path: Path) -> None:
    """A segunda fonte de relógio: o `openpyxl` sobrescreve `modified` no `save()`."""
    alvo = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(alvo, "w") as arquivo:
        arquivo.writestr(
            P.NOME_DAS_PROPRIEDADES_DO_XLSX,
            '<cp:coreProperties><dcterms:modified xsi:type="dcterms:W3CDTF">'
            "2026-10-11T13:32:05Z</dcterms:modified></cp:coreProperties>",
        )
    P._rezipar_deterministico(alvo)
    with zipfile.ZipFile(alvo) as arquivo:
        assert {info.date_time for info in arquivo.infolist()} == {P.DATA_FIXA_NO_ZIP}
        corpo = arquivo.read(P.NOME_DAS_PROPRIEDADES_DO_XLSX).decode("utf-8")
    assert P.CARIMBO_FIXO_NO_XLSX in corpo
    assert "13:32:05" not in corpo


# ═════════════════════════════════════════════════════════════════════════════
# 4 — O smoke detecta
# ═════════════════════════════════════════════════════════════════════════════


def test_smoke_detecta_artefato_proibido_injetado(tmp_path: Path) -> None:
    """O cenário que D-af existe para impedir: o Parquet proibido **dentro** do pacote."""
    pacote = tmp_path / "pacote"
    (pacote / "dados").mkdir(parents=True)
    (pacote / "dados" / "internacional_pesos.parquet").write_bytes(b"fixture sintetica")
    proibidos = {"internacional_pesos.parquet": "fixture sintetica"}
    assert P.artefatos_proibidos_no_pacote(pacote, proibidos) == (
        "dados/internacional_pesos.parquet",
    )


def test_smoke_pega_o_proibido_em_qualquer_subpasta(tmp_path: Path) -> None:
    """A comparação é pelo **nome do arquivo**: o pacote reorganiza as pastas."""
    pacote = tmp_path / "pacote"
    (pacote / "figuras").mkdir(parents=True)
    (pacote / "internacional_pesos.parquet").write_bytes(b"x")
    (pacote / "figuras" / "internacional_pesos.parquet").write_bytes(b"x")
    achados = P.artefatos_proibidos_no_pacote(pacote, {"internacional_pesos.parquet": "m"})
    assert achados == ("figuras/internacional_pesos.parquet", "internacional_pesos.parquet")


def test_smoke_nao_acusa_pacote_limpo(tmp_path: Path) -> None:
    pacote = tmp_path / "pacote"
    (pacote / "dados").mkdir(parents=True)
    (pacote / "dados" / "macro_quarterly.parquet").write_bytes(b"x")
    assert P.artefatos_proibidos_no_pacote(pacote, {"internacional_pesos.parquet": "m"}) == ()


def test_smoke_com_artefato_proibido_reprova_o_pacote_inteiro(tmp_path: Path) -> None:
    """Ponta a ponta do veredito: `rc != 0` e a falha **nomeia** o arquivo proibido."""
    pacote = tmp_path / "pacote"
    (pacote / "dados").mkdir(parents=True)
    (pacote / "dados" / "internacional_pesos.parquet").write_bytes(b"fixture sintetica")
    manifesto = manifesto_sintetico()
    manifesto["artefatos"].append(
        {
            "arquivo": "internacional_pesos.parquet",
            "publicavel": False,
            "motivo_nao_publicavel": "fixture sintetica: fora da v1",
            "sha256": SHA_SINTETICO,
        }
    )
    resultado = P.smoke(pacote, raiz=tmp_path, manifesto=manifesto)
    assert resultado.rc != 0
    nomes = [c.nome for c in resultado.falhas]
    assert any("exclusao" in nome for nome in nomes)
    evidencias = " ".join(c.evidencia for c in resultado.falhas)
    assert "internacional_pesos.parquet" in evidencias


def test_smoke_de_pacote_inexistente_reprova(tmp_path: Path) -> None:
    resultado = P.smoke(tmp_path / "nao-existe", raiz=tmp_path, manifesto=manifesto_sintetico())
    assert resultado.rc != 0
    assert resultado.falhas[0].nome.startswith("pacote:")


def test_smoke_reprova_quando_o_parquet_do_pacote_nao_bate_com_o_manifesto(
    tmp_path: Path, quadro_sintetico: pd.DataFrame
) -> None:
    """O detector de pacote defasado: o `sha256` e o nº de linhas vêm do manifesto."""
    pacote = tmp_path / "pacote"
    (pacote / "dados").mkdir(parents=True)
    quadro_sintetico.to_parquet(pacote / "dados" / "macro_quarterly.parquet")
    P.gravar_csv_ptbr(quadro_sintetico, pacote / "dados" / "macro_quarterly.csv")
    resultado = P.smoke(pacote, raiz=tmp_path, manifesto=manifesto_sintetico(linhas=83))
    nomes = {c.nome for c in resultado.falhas}
    assert "macro_quarterly: linhas Parquet == MANIFEST" in nomes
    assert "macro_quarterly: sha256 do Parquet == MANIFEST" in nomes
    assert resultado.rc != 0


def test_checagem_info_nao_reprova() -> None:
    """`INFO` é informação, não falha — o `FROZEN.txt` ausente não pode travar o smoke."""
    resultado = P.ResultadoSmoke(pacote=Path("fake"))
    resultado.informar("nota", "sem consequencia")
    assert resultado.rc == 0
    resultado.registrar("de verdade", False, "quebrou")
    assert resultado.rc == 1


# ═════════════════════════════════════════════════════════════════════════════
# 5 — O gate vem do disco (CA-090), nunca de texto digitado
# ═════════════════════════════════════════════════════════════════════════════

CAMINHO_ADR = Path("docs/decisions/adr/ADR-001-lente-de-regime-oficial.md")


def _escrever_adr(raiz: Path, status: str) -> None:
    alvo = raiz / CAMINHO_ADR
    alvo.parent.mkdir(parents=True, exist_ok=True)
    alvo.write_text(f"# ADR-001 (fixture sintetica)\n\n- **Status:** {status}\n", encoding="utf-8")


def test_gate_nasce_proposta_quando_o_adr_em_disco_esta_em_proposta(tmp_path: Path) -> None:
    _escrever_adr(tmp_path, "Proposta")
    gate = P.ler_gate(raiz=tmp_path)
    assert gate.status_adr_001 == "PROPOSTA"
    assert gate.congelada is False
    assert gate.veredito == P.VEREDITO_BLOQUEADO
    assert gate.selo == P.TEXTO_DO_GATE_NAO_PUBLICAVEL


def test_gate_libera_quando_o_adr_em_disco_esta_aceita(tmp_path: Path) -> None:
    """O dia em que o Edu aceitar o ADR-001, **nenhuma linha de código** muda: o selo cai."""
    _escrever_adr(tmp_path, "Aceita")
    gate = P.ler_gate(raiz=tmp_path)
    assert gate.congelada is True
    assert gate.selo == ""
    assert gate.veredito == P.VEREDITO_LIBERADO_PELO_ADR


def test_gate_falha_fechada_sem_adr_em_disco(tmp_path: Path) -> None:
    """ADR ausente ⇒ `DESCONHECIDO` e **não** congelado. Nunca o contrário."""
    gate = P.ler_gate(raiz=tmp_path)
    assert gate.congelada is False
    assert gate.status_adr_001 == "DESCONHECIDO"


def test_o_readme_carrega_o_selo_lido_do_disco(tmp_path: Path) -> None:
    """CA-115: o selo do README nasce do ADR em disco, não de uma frase no fonte."""
    _escrever_adr(tmp_path, "Proposta")
    texto = P.markdown_do_readme(
        gate=P.ler_gate(raiz=tmp_path),
        manifesto=manifesto_sintetico(),
        frozen_presente=False,
    )
    assert P.MARCADOR_DO_GATE in texto.split("## 1.")[0]
    assert "PROPOSTA" in texto
    assert P.VERSAO_DO_PACOTE in texto
    assert SHA_SINTETICO in texto


def test_o_readme_perde_o_selo_quando_o_adr_e_aceito(tmp_path: Path) -> None:
    _escrever_adr(tmp_path, "Aceita")
    texto = P.markdown_do_readme(
        gate=P.ler_gate(raiz=tmp_path),
        manifesto=manifesto_sintetico(),
        frozen_presente=True,
    )
    assert P.MARCADOR_DO_GATE not in texto.split("## 1.")[0]
    assert P.NOME_FROZEN in texto


def test_o_readme_e_descritivo_nos_dois_estados_do_gate(tmp_path: Path) -> None:
    """Anti-SPEC §10.2.10 — a guarda de linguagem roda sobre o texto gerado, não sobre o fonte."""
    from apps.analysis import _linguagem

    for status in ("Proposta", "Aceita"):
        _escrever_adr(tmp_path, status)
        texto = P.markdown_do_readme(
            gate=P.ler_gate(raiz=tmp_path),
            manifesto=manifesto_sintetico(),
            frozen_presente=False,
        )
        _linguagem.conferir_linguagem_descritiva(texto, origem=f"README ({status})")


def test_o_readme_e_deterministico(tmp_path: Path) -> None:
    _escrever_adr(tmp_path, "Proposta")
    gate = P.ler_gate(raiz=tmp_path)
    manifesto = manifesto_sintetico()
    primeiro = P.markdown_do_readme(gate=gate, manifesto=manifesto, frozen_presente=False)
    segundo = P.markdown_do_readme(gate=gate, manifesto=manifesto, frozen_presente=False)
    assert primeiro == segundo


def test_o_readme_nomeia_o_artefato_excluido_e_o_motivo(tmp_path: Path) -> None:
    """A exclusão é **declarada** ao gestor, não só ausente do diretório."""
    _escrever_adr(tmp_path, "Proposta")
    manifesto = manifesto_sintetico()
    manifesto["artefatos"].append(
        {
            "arquivo": "internacional_pesos.parquet",
            "publicavel": False,
            "motivo_nao_publicavel": "fixture sintetica: fora da v1",
            "sha256": SHA_SINTETICO,
        }
    )
    texto = P.markdown_do_readme(
        gate=P.ler_gate(raiz=tmp_path), manifesto=manifesto, frozen_presente=False
    )
    assert "internacional_pesos.parquet" in texto
    assert "fixture sintetica: fora da v1" in texto


def test_hashes_txt_e_ascii_puro_e_nao_se_autoinclui(tmp_path: Path) -> None:
    """Anti-SPEC §10.4.26: o `HASHES.txt` é lido no terminal e no Notepad do Windows."""
    _escrever_adr(tmp_path, "Proposta")
    pacote = tmp_path / "pacote"
    pacote.mkdir()
    (pacote / "a.txt").write_text("a", encoding="utf-8")
    registros = P.gravar_hashes(
        pacote, manifesto=manifesto_sintetico(), gate=P.ler_gate(raiz=tmp_path)
    )
    texto = (pacote / P.NOME_HASHES).read_text(encoding="utf-8")
    assert [nome for nome, _, _ in registros] == ["a.txt"]
    assert texto.isascii()
    assert "DATASET_SHA256" in texto
    assert f"  {P.NOME_HASHES}" not in texto


def test_limpar_apaga_o_pacote_e_o_zip(tmp_path: Path) -> None:
    """`--limpar` existe por causa da interação com `verify` — ver `AVISO_SOBRE_VERIFY`."""
    pacote = tmp_path / "data" / "canonical" / "entrega" / P.NOME_DO_PACOTE
    (pacote / "dados").mkdir(parents=True)
    (pacote / "dados" / "x.txt").write_text("x", encoding="utf-8")
    arquivo_zip = pacote.with_name(f"{pacote.name}.zip")
    arquivo_zip.write_bytes(b"PK")
    removidos = P.limpar(raiz=tmp_path)
    assert set(removidos) == {pacote, arquivo_zip}
    assert not pacote.exists()
    assert not arquivo_zip.exists()


def test_limpar_e_idempotente_quando_nao_ha_pacote(tmp_path: Path) -> None:
    assert P.limpar(raiz=tmp_path) == ()


def test_o_aviso_sobre_verify_nomeia_a_ordem_e_a_saida() -> None:
    """A interação com F-029 é **declarada** na saída do comando, não só no runbook."""
    texto = " ".join(P.AVISO_SOBRE_VERIFY)
    assert "verify" in texto
    assert "freeze" in texto
    assert "--limpar" in texto
    assert texto.isascii()


# ═════════════════════════════════════════════════════════════════════════════
# 6 — Os textos vêm do disco: 10 limitações, 12 decisões, substituições
# ═════════════════════════════════════════════════════════════════════════════


def test_itens_numerados_colapsa_continuacao_indentada() -> None:
    markdown = "1. primeira\n   continua aqui\n2. segunda\n\ntexto fora da lista\n"
    assert P._itens_numerados(markdown) == ("primeira continua aqui", "segunda")


def test_limitacoes_bloqueia_quando_o_caderno_nao_tem_a_celula(tmp_path: Path) -> None:
    alvo = tmp_path / P.CAMINHO_NOTEBOOK_LIMPO
    alvo.parent.mkdir(parents=True)
    alvo.write_text(json.dumps({"cells": [{"id": "outra", "source": []}]}), encoding="utf-8")
    with pytest.raises(SerieInconsistente, match="FONTE INDISPONÍVEL"):
        P.limitacoes_do_notebook(raiz=tmp_path)


def test_limitacoes_le_a_celula_do_caderno(tmp_path: Path) -> None:
    alvo = tmp_path / P.CAMINHO_NOTEBOOK_LIMPO
    alvo.parent.mkdir(parents=True)
    alvo.write_text(
        json.dumps(
            {
                "cells": [
                    {
                        "id": P._ID_DA_CELULA_DAS_LIMITACOES,
                        "source": ["## 8. Limitacoes\n", "1. uma\n", "2. duas\n"],
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    assert P.limitacoes_do_notebook(raiz=tmp_path) == ("uma", "duas")


def test_decisoes_bloqueia_quando_a_contagem_nao_e_12(tmp_path: Path) -> None:
    alvo = tmp_path / P.CAMINHO_LIMITACOES
    alvo.parent.mkdir(parents=True)
    alvo.write_text(
        "## 7. Decisoes\n\n| # | a | b | c |\n|---|---|---|---|\n| D1 | x | y | z |\n",
        encoding="utf-8",
    )
    with pytest.raises(SerieInconsistente, match="FONTE INDISPONÍVEL"):
        P.decisoes_a_ratificar(raiz=tmp_path)


def test_substituicoes_bloqueia_sem_a_secao_1_do_decisions_log(tmp_path: Path) -> None:
    """CA-117 sem o texto da §1 é `BLOQUEADO` — o pacote não sai sem a comunicação escrita."""
    alvo = tmp_path / P.CAMINHO_DECISIONS_LOG
    alvo.parent.mkdir(parents=True)
    alvo.write_text("# DECISIONS_LOG\n\n## 2. Outra coisa\n", encoding="utf-8")
    with pytest.raises(SerieInconsistente, match="FONTE INDISPONÍVEL"):
        P.texto_das_substituicoes(raiz=tmp_path)


def test_substituicoes_recorta_so_a_secao_1(tmp_path: Path) -> None:
    alvo = tmp_path / P.CAMINHO_DECISIONS_LOG
    alvo.parent.mkdir(parents=True)
    alvo.write_text(
        "# DECISIONS_LOG\n\n## 1. Substituicoes de fonte\n\nS-01 fixture\n\n## 2. Outra\n\nnao\n",
        encoding="utf-8",
    )
    texto = P.texto_das_substituicoes(raiz=tmp_path)
    assert "S-01 fixture" in texto
    assert "## 2. Outra" not in texto
    assert "nao" not in texto


def test_sem_markdown_limpa_enfase_e_crase() -> None:
    assert P._sem_markdown("**D1** — `V1`<br>oficial") == "D1 — V1 oficial"


# ═════════════════════════════════════════════════════════════════════════════
# 7 — Acervo real (padrão `Acervo`; desligado por LDC_IGNORAR_DATA=1)
# ═════════════════════════════════════════════════════════════════════════════


def ignorar_data_real() -> bool:
    """`LDC_IGNORAR_DATA` ligada? (vazio, `0`, `false` e `nao` contam como desligada)."""
    return os.environ.get(VAR_IGNORAR_DATA, "").strip().lower() not in {
        "",
        "0",
        "false",
        "nao",
        "no",
    }


def canonical_disponivel() -> bool:
    """Tudo o que a montagem real precisa está em disco nesta máquina?"""
    alvos = [P.NOME_MANIFESTO, P.NOME_HEATMAP_DADOS, P.CAMINHO_NOTEBOOK_HTML.as_posix()]
    alvos += [f"{nome}.parquet" for nome in P.TABELAS_DO_PACOTE]
    return all(caminho_canonical(nome).is_file() for nome in alvos)


acervo_real = pytest.mark.skipif(
    ignorar_data_real() or not canonical_disponivel(),
    reason=f"data/canonical/ ausente ou {VAR_IGNORAR_DATA} ligada — {COMO_GERAR}",
)


@pytest.fixture(scope="module")
def pacote_real():  # noqa: ANN201 - ResultadoMontagem
    """O pacote real, montado **no destino padrão** (dentro de `data/`, via script — CA-114)."""
    if ignorar_data_real() or not canonical_disponivel():
        pytest.skip("acervo real indisponível")
    return P.montar()


@acervo_real
def test_acervo_o_gate_em_disco_continua_fechado(pacote_real) -> None:  # noqa: ANN001
    """Enquanto ADR-001 estiver em `Proposta`, o pacote nasce selado. Se este teste falhar,
    o ADR mudou de status e o fluxo de publicação precisa de revisão humana — não de código."""
    assert pacote_real.gate.congelada is False
    assert pacote_real.gate.status_adr_001 == "PROPOSTA"
    assert pacote_real.gate.veredito == P.VEREDITO_BLOQUEADO


@acervo_real
def test_acervo_o_pacote_tem_todos_os_itens_de_ca113(pacote_real) -> None:  # noqa: ANN001
    """CA-113: as tabelas nos dois formatos, as figuras, as 1-páginas e o README com hash."""
    pacote = pacote_real.pacote
    for nome in P.TABELAS_DO_PACOTE:
        assert (pacote / "dados" / f"{nome}.parquet").is_file()
        assert (pacote / "dados" / f"{nome}.csv").is_file()
    for nome in P.FIGURAS_DO_PACOTE:
        for extensao in P.EXTENSOES_DE_FIGURA:
            assert (pacote / "figuras" / f"{nome}.{extensao}").is_file()
    for nome in P.TABELAS_DE_UMA_PAGINA:
        for extensao in P.EXTENSOES_DE_TABELA:
            assert (pacote / "tabelas" / f"{nome}.{extensao}").is_file()
    for nome in (P.NOME_README, P.NOME_LIMITACOES, P.NOME_SUBSTITUICOES, P.NOME_XLSX):
        assert (pacote / nome).is_file()
    assert (pacote / "notebook" / P.CAMINHO_NOTEBOOK_HTML.name).is_file()
    assert (pacote / "notebook" / P.CAMINHO_NOTEBOOK_LIMPO.name).is_file()


@acervo_real
def test_acervo_nenhum_artefato_nao_publicavel_entrou(pacote_real) -> None:  # noqa: ANN001
    """D-af no acervo real: hoje `internacional_pesos.parquet` é o único proibido."""
    assert set(pacote_real.excluidos) == {"internacional_pesos.parquet"}
    assert P.artefatos_proibidos_no_pacote(pacote_real.pacote, pacote_real.excluidos) == ()


@acervo_real
def test_acervo_o_parquet_do_pacote_e_copia_byte_a_byte_do_canonico(
    pacote_real,  # noqa: ANN001
) -> None:
    """O `sha256` do Parquet no pacote **é** o do `MANIFEST.json` — nada foi reserializado."""
    indice = P.artefatos_do_manifesto(pacote_real.manifesto)
    for nome in P.TABELAS_DO_PACOTE:
        no_pacote = P.sha256_do_arquivo(pacote_real.pacote / "dados" / f"{nome}.parquet")
        assert no_pacote == indice[f"{nome}.parquet"]["sha256"], nome


@acervo_real
def test_acervo_duas_montagens_dao_o_mesmo_sha256_do_zip(pacote_real) -> None:  # noqa: ANN001
    """CA-121 aplicado ao pacote: idempotência ponta a ponta, com o XLSX dentro."""
    segunda = P.montar()
    assert segunda.sha256_do_zip == pacote_real.sha256_do_zip
    assert segunda.sha256_do_zip == P.sha256_do_arquivo(segunda.zip)


@acervo_real
def test_acervo_o_smoke_aprova_o_pacote_recem_montado(pacote_real) -> None:  # noqa: ANN001
    resultado = P.smoke(pacote_real.pacote, manifesto=pacote_real.manifesto)
    assert resultado.rc == 0, [(c.nome, c.evidencia) for c in resultado.falhas]


@acervo_real
def test_acervo_o_xlsx_tem_as_8_abas_e_zero_formula(pacote_real) -> None:  # noqa: ANN001
    """ "Valores colados": nenhuma célula começa com `=`, então nada recalcula ao abrir."""
    from openpyxl import load_workbook

    livro = load_workbook(pacote_real.pacote / P.NOME_XLSX, read_only=True)
    try:
        assert tuple(livro.sheetnames) == P.ABAS_ESPERADAS
        formulas = [
            f"{aba}!{celula.coordinate}"
            for aba in livro.sheetnames
            for linha in livro[aba].iter_rows()
            for celula in linha
            if isinstance(celula.value, str) and celula.value.startswith("=")
        ]
        assert formulas == []
    finally:
        livro.close()


@acervo_real
def test_acervo_a_matriz_do_heatmap_e_23x5_com_n_ao_lado(pacote_real) -> None:  # noqa: ANN001
    del pacote_real
    heatmap = pd.read_parquet(caminho_canonical(P.NOME_HEATMAP_DADOS))
    matriz, ativos, regimes = P.matriz_do_heatmap(heatmap)
    assert len(ativos) == P.N_ATIVOS_NO_HEATMAP
    assert len(regimes) == P.N_REGIMES_NO_HEATMAP
    assert matriz.shape == (P.N_ATIVOS_NO_HEATMAP, 1 + 2 * P.N_REGIMES_NO_HEATMAP)
    for regime in regimes:
        assert f"n ({regime})" in matriz.columns
        assert matriz[f"n ({regime})"].notna().all()


@acervo_real
def test_acervo_as_10_limitacoes_e_as_12_decisoes_vem_do_disco() -> None:
    """A contagem é trava: se a página ou o caderno mudar, o pacote avisa em vez de calar."""
    assert len(P.limitacoes_do_notebook()) == P.N_LIMITACOES
    decisoes = P.decisoes_a_ratificar()
    assert len(decisoes) == P.N_DECISOES
    assert {d["id"] for d in decisoes} >= {"D1", "D3", "D5", "D6"}
    assert all(d["padrao_provisorio"] for d in decisoes)


@acervo_real
def test_acervo_o_texto_das_substituicoes_tem_as_12_e_os_4_campos() -> None:
    """CA-117: fonte pedida, fonte usada, por quê e impacto — para cada substituição."""
    texto = P.texto_das_substituicoes()
    for indice in range(1, 13):
        assert f"S-{indice:02d}" in texto
    for campo in ("Fonte pedida", "Fonte usada", "Por quê", "Impacto na comparabilidade"):
        assert campo in texto


@acervo_real
def test_acervo_o_hashes_cobre_todo_arquivo_do_pacote(pacote_real) -> None:  # noqa: ANN001
    """CA-114 verificável: o `HASHES.txt` fecha com o conteúdo do diretório, arquivo por arquivo."""
    texto = (pacote_real.pacote / P.NOME_HASHES).read_text(encoding="utf-8")
    for nome, _, sha in pacote_real.arquivos:
        assert f"{sha}" in texto
        assert nome in texto


@acervo_real
def test_acervo_o_zip_tem_um_arquivo_a_mais_que_a_lista_o_hashes(
    pacote_real,  # noqa: ANN001
) -> None:
    """O `HASHES.txt` não se autoinclui na lista, mas **vai** dentro do `.zip`."""
    with zipfile.ZipFile(pacote_real.zip) as arquivo:
        nomes = {info.filename for info in arquivo.infolist()}
    assert len(nomes) == len(pacote_real.arquivos) + 1
    assert f"{pacote_real.pacote.name}/{P.NOME_HASHES}" in nomes
